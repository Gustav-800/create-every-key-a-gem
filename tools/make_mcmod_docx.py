from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from lxml import html


ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "docs" / "mcmod-body-richtext.html"
OUT = ROOT / "docs" / "mcmod-body-richtext.docx"


def set_font(run, size=11, bold=None, color=None):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_inline(paragraph, element):
    if element.text:
        run = paragraph.add_run(element.text)
        set_font(run)
    for child in element:
        if child.tag == "strong":
            run = paragraph.add_run(child.text or "")
            set_font(run, bold=True)
        elif child.tag == "a":
            run = paragraph.add_run(child.text or child.get("href", ""))
            set_font(run, color=(31, 111, 235))
        elif child.tag == "br":
            paragraph.add_run().add_break()
        else:
            add_inline(paragraph, child)
        if child.tail:
            run = paragraph.add_run(child.tail)
            set_font(run)


def main() -> None:
    tree = html.fromstring(SOURCE.read_text(encoding="utf-8"))
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)

    body = tree.body
    for element in body:
        tag = element.tag.lower() if isinstance(element.tag, str) else ""
        if tag == "h1":
            heading = doc.add_heading(level=1)
            run = heading.add_run(element.text_content().strip())
            set_font(run, size=20, bold=True, color=(31, 78, 121))
        elif tag == "h2":
            heading = doc.add_heading(level=2)
            run = heading.add_run(element.text_content().strip())
            set_font(run, size=14, bold=True, color=(31, 78, 121))
        elif tag == "p":
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(6)
            add_inline(paragraph, element)
            if "meta" in (element.get("class") or ""):
                for run in paragraph.runs:
                    run.font.size = Pt(9)
                    run.font.color.rgb = RGBColor(102, 102, 102)
        elif tag in ("ul", "ol"):
            style = "List Bullet" if tag == "ul" else "List Number"
            for li in element.findall("li"):
                paragraph = doc.add_paragraph(style=style)
                paragraph.paragraph_format.space_after = Pt(3)
                add_inline(paragraph, li)

    doc.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
